"""Local, transactional imports. No network clients or automatic source-folder scans."""
import csv
import hashlib
import io
import json
import math
import os
import re
import shutil
import uuid
from datetime import date, datetime
from pathlib import Path
import openpyxl

from analytics_data import extract_analytics
from port_parser import extract

APP = Path(__file__).resolve().parent
REPORT_NAMES = [f'Attribution Analysis 260{m}.xlsx' for m in range(4, 9)] + [
    'Attribution Analysis MTD.xlsx',
    'Risk Analysis JIKA.xlsx', 'Risk Analysis JPY.xlsx', 'Risk Analysis USD.xlsx']
SELF_START_DATE = '2026-03-31'
TRADE_FIELDS = ['trade_id', 'date', 'ticker', 'event_type', 'quantity_delta',
                'amount_jpy', 'amount_basis', 'pair_id', 'notes', 'status']
TRADE_HEADERS = ['取引ID', '取引日', '銘柄コード', '取引区分', '数量増減',
                 '売買金額（円）', '金額区分', '関連取引ID', '備考', '確認状態']
EVENT_JA = {'買付': 'buy', '売却': 'sell', '分割': 'split', '入金': 'cash_in', '出金': 'cash_out', 'その他': 'other'}
STATUS_JA = {'確認済み': 'confirmed', '未確認': 'pending', '取消': 'void'}
BASIS_JA = {'実額': 'actual', '概算': 'estimated'}


def dump(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def csv_bytes(rows, fields=TRADE_FIELDS):
    out = io.StringIO(newline='')
    writer = csv.DictWriter(out, fieldnames=fields, extrasaction='ignore')
    writer.writeheader()
    for row in rows:
        # Text exported for Excel must never become a formula.
        writer.writerow({k: "'" + v if isinstance(v, str) and v.startswith(('=', '+', '@', '\t', '\r'))
                         else v for k, v in row.items()})
    return out.getvalue().encode('utf-8-sig')


def read_trades(payload):
    reader = csv.DictReader(io.StringIO(payload.decode('utf-8-sig')))
    if reader.fieldnames != TRADE_FIELDS:
        raise ValueError('売買CSVの列名が一致しません。ダウンロードしたテンプレートをお使いください。')
    rows, ids = [], set()
    for n, row in enumerate(reader, 2):
        row = {k: (v or '').strip() for k, v in row.items()}
        if not any(row.values()):
            continue
        if not row['trade_id'] or row['trade_id'] in ids:
            raise ValueError(f'{n}行目：trade_idが未入力、または重複しています。')
        ids.add(row['trade_id'])
        if row['date']:
            date.fromisoformat(row['date'])
        elif row['status'] == 'confirmed':
            raise ValueError('確認済みの取引には取引日が必要です。')
        if not row['ticker'] or row['status'] not in ('confirmed', 'pending', 'void'):
            raise ValueError(f'{n}行目：tickerを入力し、statusにconfirmed / pending / voidを指定してください。')
        if row['event_type'] not in ('buy', 'sell', 'split', 'cash_in', 'cash_out', 'other') and not (row['status'] != 'confirmed' and not row['event_type']):
            raise ValueError(f'{n}行目：event_typeにbuy / sell / split / cash_in / cash_out / otherを指定してください。')
        for field in ('quantity_delta', 'amount_jpy'):
            if row[field]:
                row[field] = float(row[field])
                if not math.isfinite(row[field]):
                    raise ValueError(f'{n}行目：有効な数値を入力してください。')
        if row['status'] == 'confirmed' and row['event_type'] in ('buy', 'sell'):
            if row['amount_jpy'] == '' or row['amount_basis'] not in ('actual', 'estimated'):
                raise ValueError(f'{n}行目：売買の確認には円建て金額とactual / estimatedの指定が必要です。')
            if (row['event_type'] == 'buy' and row['amount_jpy'] <= 0) or (row['event_type'] == 'sell' and row['amount_jpy'] >= 0):
                raise ValueError(f'{n}行目：買付金額は正、売却金額は負で入力してください。')
        rows.append(row)
    return rows


def read_trade_workbook(payload):
    """Accept the maintained ledger, including rows appended outside its Excel table."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(payload), data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError('売買台帳をExcel形式（.xlsx）で選択してください。') from exc
    try:
        if '売買台帳' not in wb.sheetnames:
            raise ValueError('「売買台帳」シートが見つかりません。配布したExcelを使用してください。')
        sheet = wb['売買台帳']
        header = None
        data = []
        for n, cells in enumerate(sheet.iter_rows(values_only=True), 1):
            metadata = list(cells[10:14])
            cells = list(cells[:10]) + [None] * max(0, 10 - len(cells))
            if header is None:
                if cells == TRADE_HEADERS or cells == TRADE_FIELDS:
                    header = n
                elif n >= 20:
                    raise ValueError('売買台帳の列見出しが一致しません。列名を変更しないでください。')
                continue
            if not any(v is not None and v != '' for v in cells):
                continue
            row = dict(zip(TRADE_FIELDS, cells))
            if isinstance(row['date'], (date, datetime)):
                row['date'] = row['date'].strftime('%Y-%m-%d')
            row['event_type'] = EVENT_JA.get(row['event_type'], row['event_type'])
            row['status'] = STATUS_JA.get(row['status'], row['status'])
            row['amount_basis'] = BASIS_JA.get(row['amount_basis'], row['amount_basis'])
            for key in TRADE_FIELDS:
                if row[key] is None:
                    row[key] = ''
            try:
                # Reuse the canonical rules without applying CSV export sanitization.
                out = io.StringIO()
                writer = csv.DictWriter(out, fieldnames=TRADE_FIELDS)
                writer.writeheader()
                writer.writerow(row)
                normalized = read_trades(out.getvalue().encode('utf-8'))[0]
            except (ValueError, IndexError) as exc:
                raise ValueError(f'Excel {n}行目の入力を確認してください。日付はYYYY-MM-DD、金額は数値で入力します。詳細：{exc}') from exc
            for key, value in zip(('_detection_id', '_detected_start', '_detected_end', '_required'), metadata):
                if value is not None:
                    normalized[key] = value.strftime('%Y-%m-%d') if isinstance(value, (date, datetime)) else str(value)
            data.append(normalized)
        if header is None:
            raise ValueError('売買台帳の列見出しが見つかりません。')
        ids = [row['trade_id'] for row in data]
        if len(ids) != len(set(ids)):
            raise ValueError('取引IDが重複しています。1取引につき1行・1IDにしてください。')
        return sorted(data, key=lambda r: (r['date'], r['trade_id']))
    finally:
        wb.close()


def compare_trades(previous, incoming):
    old = {r['trade_id']: r for r in previous}
    new = {r['trade_id']: r for r in incoming}
    result = dict(added=[], modified=[], replaced=[], cancelled=[], missing=[], unchanged=0,
                  pending=sum(r['status'] == 'pending' for r in incoming), total=len(incoming))
    for key, row in new.items():
        before = old.get(key)
        item = dict(id=key, before=before, after=row,
                    fields=[TRADE_HEADERS[i] for i,k in enumerate(TRADE_FIELDS) if before and before[k] != row[k]])
        if before is None:
            result['added'].append(item)
        elif before == row:
            result['unchanged'] += 1
        elif row['status'] == 'void' and before['status'] != 'void':
            result['cancelled'].append(item)
        else:
            result['modified'].append(item)
    result['missing'] = [old[key] for key in old.keys() - new.keys()]
    # API-generated rows may replace deleted legacy IDs. Require a unique match
    # in both directions; similar or ambiguous trades remain separate for review.
    def same_trade(before, after):
        if before['status'] == 'void' or after['status'] == 'void':
            return False
        if not after.get('_detection_id') or not after['trade_id'].startswith('API-'):
            return False
        if before['ticker'].strip().upper() != after['ticker'].strip().upper() or before['event_type'] != after['event_type']:
            return False
        if not before.get('date') or not after.get('date') or before['quantity_delta'] == '' or after['quantity_delta'] == '':
            return False
        return (abs((date.fromisoformat(before['date'])-date.fromisoformat(after['date'])).days) <= 3
                and not math.isclose(float(before['quantity_delta']), 0, abs_tol=1e-7)
                and math.isclose(float(before['quantity_delta']), float(after['quantity_delta']), rel_tol=0, abs_tol=1e-6))
    pairs=[(before,item) for before in result['missing'] for item in result['added'] if same_trade(before,item['after'])]
    replaced_old=set();replaced_new=set()
    for before,item in pairs:
        if sum(b['trade_id']==before['trade_id'] for b,_ in pairs)!=1 or sum(i['id']==item['id'] for _,i in pairs)!=1:
            continue
        result['replaced'].append(dict(item,before=before,previous_id=before['trade_id'],
            fields=[TRADE_HEADERS[i] for i,k in enumerate(TRADE_FIELDS) if before[k]!=item['after'][k]]))
        replaced_old.add(before['trade_id']);replaced_new.add(item['id'])
    result['added']=[i for i in result['added'] if i['id'] not in replaced_new]
    result['missing']=[r for r in result['missing'] if r['trade_id'] not in replaced_old]
    result['can_calculate'] = not result['missing']
    return result


def trade_payload(rows):
    confirmed = [r for r in rows if r['status'] == 'confirmed']
    trades = [dict(id=r['trade_id'], date=r['date'], security=r['ticker'],
                   label=r['event_type'], delta=r['quantity_delta'] or 0,
                   estimate=r['amount_jpy'] if r['event_type'] in ('buy', 'sell') else None,
                   pair=r['pair_id'], note=r['notes'], amount_basis=r['amount_basis']) for r in confirmed]
    dates = [r['date'] for r in confirmed]
    return dict(rows=trades, start=min(dates, default=''), end=max(dates, default=''),
                coverage_note='確認済み台帳に記載された取引のみ集計。未登録日は売買ゼロを意味しません。' if dates else
                '売買記録は未登録です。取込・通知画面から売買台帳のExcelを選択して計算してください。')


def build(folder, trades):
    reports = [r for name in REPORT_NAMES[:6] for r in extract(folder / name)]
    if len(reports) != 6:
        raise ValueError('各ファイルには要因分析レポートを1件だけ含めてください。')
    first = reports[0]
    for r in reports:
        if r['portfolio'] != first['portfolio'] or r['classification'] != first['classification']:
            raise ValueError('要因分析レポート間でポートフォリオ・ベンチマーク・分類が一致していません。')
        if 'JPY' not in r['currency'] or r['start'] >= r['end']:
            raise ValueError('要因分析レポートはJPY建てで、有効な開始日・終了日を設定してください。')
        if any(not isinstance(r['total'].get(k), (int, float)) for k in ('rp', 'rb', 'allocation', 'selection')):
            raise ValueError('要因分析レポートに合計収益率または要因効果の項目がありません。')
        if not any(not g.get('no_summary') for s in r['sectors'] for g in s['groups']):
            raise ValueError('要因分析レポートに業種グループ（Industry Group）の階層がありません。')
    for r, name in zip(reports[:5], REPORT_NAMES[:5]):
        suffix = name.split()[-1][:4]
        if r['end'][:7] != f'20{suffix[:2]}-{suffix[2:]}':
            raise ValueError(name + '：対象月がファイル名と一致していません。')
    mtd = next(r for r in reports if r['kind'] == 'MTD')
    portfolio, sep, benchmark = first['portfolio'].partition(' vs. ')
    if not sep:
        raise ValueError('レポート内のポートフォリオ名とベンチマーク名を読み取れません。')
    analytics = extract_analytics(folder, SELF_START_DATE, portfolio, benchmark)
    if mtd['end'] != analytics['end']:
        raise ValueError(f'MTD・JPY日次データの終了日を揃えてください。日次データの終了日は{analytics["end"]}です。')
    if any(r['start'] < analytics['start'] or r['end'] > analytics['end'] for r in reports):
        raise ValueError('日次データが要因分析レポートの全期間をカバーしていません。')
    analytics['comparison_reports'] = reports
    reports.sort(key=lambda r: r['end'])
    page = (APP / 'report_template.html').read_text(encoding='utf-8-sig')
    for marker, obj in [('__PORT_DATA__', reports), ('__TRADE_DATA__', trade_payload(trades)), ('__ANALYTICS_DATA__', analytics)]:
        page = page.replace(marker, dump(obj).replace('<', '\\u003c').replace('>', '\\u003e').replace('&', '\\u0026'))
    page = page.replace('__ANALYTICS_SCRIPT__', (APP / 'analytics.js').read_text(encoding='utf-8-sig'))
    page = page.replace('__DOWNLOAD_URL__', '/download/' + folder.name + '.html')
    (folder / 'report.html').write_text(page, encoding='utf-8')
    return reports, analytics


class Store:
    def __init__(self, root):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)
        self.state_path = self.root / 'state.json'

    def state(self):
        return json.loads(self.state_path.read_text(encoding='utf-8')) if self.state_path.exists() else dict(
            run=None, trades=[], notices=[], seen_candidates=[], start=None, end=None)

    def save(self, state):
        temp = self.root / 'state.tmp'
        temp.write_text(dump(state), encoding='utf-8')
        os.replace(temp, self.state_path)

    def report(self):
        run = self.state()['run']
        return self.root / 'runs' / run / 'report.html' if run else None

    def saved_report(self, run):
        if not re.fullmatch(r'[a-f0-9]{32}', run):
            return None
        path = self.root / 'runs' / run / 'report.html'
        return path if path.is_file() else None

    def history(self):
        items = []
        for folder in (self.root / 'runs').glob('*'):
            if not self.saved_report(folder.name):
                continue
            meta = folder / 'summary.json'
            item = json.loads(meta.read_text(encoding='utf-8')) if meta.exists() else dict(
                created=datetime.fromtimestamp((folder / 'report.html').stat().st_mtime).isoformat(timespec='seconds'),
                start=None, end=None, portfolio='以前のレポート')
            items.append(dict(item, id=folder.name))
        return sorted(items, key=lambda x:x['created'], reverse=True)

    def import_reports(self, files, trades=None, checked_until=None):
        if set(files) != set(REPORT_NAMES):
            missing = set(REPORT_NAMES) - set(files)
            extra = set(files) - set(REPORT_NAMES)
            raise ValueError('9ファイルをまとめて読み込んでください。不足：' + ', '.join(sorted(missing)) + '；対象外：' + ', '.join(sorted(extra)))
        state = self.state()
        run = uuid.uuid4().hex
        folder = self.root / 'runs' / run
        folder.mkdir(parents=True)
        try:
            for name, payload in files.items():
                (folder / name).write_bytes(payload)
            rows = state['trades'] if trades is None else trades
            reports, analytics = build(folder, rows)
            if checked_until and analytics['end'] > checked_until:
                raise ValueError('APIの確認終了日をレポートの終了日以降にして、取引を再確認してください。')
            (folder / 'summary.json').write_text(dump(dict(
                created=datetime.now().isoformat(timespec='seconds'), start=analytics['start'], end=analytics['end'],
                portfolio=reports[0]['portfolio'], trades=len(rows))), encoding='utf-8')
            state.update(run=run, trades=rows,
                         start=analytics['start'], end=analytics['end'], updated=datetime.now().isoformat(timespec='seconds'))
            self.save(state)
        except Exception:
            shutil.rmtree(folder)
            raise
        return state

    def check_workbook(self, payload):
        incoming = read_trade_workbook(payload)
        state = self.state()
        result = compare_trades(state['trades'], incoming)
        # A comparison becomes stale when the workbook or saved result changes.
        fingerprint = dump([state['run'], state['trades'], incoming]).encode('utf-8')
        result['check_token'] = hashlib.sha256(fingerprint).hexdigest()
        result['first_import'] = not state['run'] and not state['trades']
        return result

    def calculate(self, files, workbook, check_token, checked_until=None):
        check = self.check_workbook(workbook)
        if check_token != check['check_token']:
            raise ValueError('台帳または保存済み結果が変わりました。「新規取引を確認」をもう一度押してください。')
        if not check['can_calculate']:
            raise ValueError('既存の取引が台帳から抜けています。行を戻すか、同じ取引IDで確認状態を「取消」にしてください。')
        rows = read_trade_workbook(workbook)
        return self.import_reports(files, rows, checked_until=checked_until)

    def import_trades(self, payload):
        incoming = read_trades(payload)
        state = self.state()
        by_id = {r['trade_id']: r for r in state['trades']}
        changed = 0
        for row in incoming:
            if by_id.get(row['trade_id']) != row:
                changed += 1
                by_id[row['trade_id']] = row
        if not changed:
            return state
        merged = sorted(by_id.values(), key=lambda r: (r['date'], r['trade_id']))
        if state['run']:
            folder = self.root / 'runs' / state['run']
            state = self.import_reports({n: (folder / n).read_bytes() for n in REPORT_NAMES}, merged)
        else:
            state['trades'] = merged
        state['notices'].append(dict(id=uuid.uuid4().hex, date=date.today().isoformat(), type='trade', read=False,
                                     text=f'売買台帳を{changed}件追加・更新しました。確認済み（confirmed）の売買金額のみ「月次推移・売買」に集計します。'))
        self.save(state)
        return state

    def mark_read(self):
        state = self.state()
        for notice in state['notices']:
            notice['read'] = True
        self.save(state)
        return state
