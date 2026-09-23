from pathlib import Path
from collections import defaultdict
import json, re, sys
import openpyxl


ROOT=Path(__file__).resolve().parent
KEYS=['wp','wb','wd','cp','cb','cd','rp','rb','rd','allocation','selection','currency','transaction','leverage','pricing','attribution']
def extract(path):
    workbook=openpyxl.load_workbook(path,data_only=True)
    reports=[]
    for sheet in workbook:
        header=str(sheet.cell(5,1).value)
        match=re.search(r'(\d{2}/\d{2}/\d{4})\s*-\s*(\d{2}/\d{2}/\d{4})',header)
        if not match:continue
        def iso(s):
            m,d,y=s.split('/');return f'{y}-{m}-{d}'
        report=dict(start=iso(match[1]),end=iso(match[2]),source=Path(path).name,
                    portfolio=str(sheet.cell(3,1).value),classification=str(sheet.cell(7,1).value),
                    currency=str(sheet.cell(6,1).value),run=str(sheet.cell(4,1).value),sectors=[])
        report['benchmark']=report['portfolio'].split(' vs. ')[-1]
        suffix=Path(path).stem.split()[-1]
        report['kind']='SPECIAL' if Path(path).stem.endswith('Special Day') else suffix if suffix in ('FYTD','MTD','WTD') else 'MONTH'
        report['title']=str(sheet.cell(1,1).value)
        expected=['Avg % Wgt']*3+['CTR']*3+['Tot Rtn']*3+['Alloc','Selec','Curr','Transact','Lev','Px Diff','Tot Attr']
        headers=[]
        for j in range(2,18):
            value=sheet.cell(9,j).value
            headers.append(value if value is not None else headers[-1] if headers else None)
        if headers!=expected:
            raise ValueError('レポートの列構成が想定と異なります：'+str(path))
        sector=None;group=None
        def record(row):
            obj={'name':sheet.cell(row,1).value,'row':row}
            obj.update({k:sheet.cell(row,j).value for j,k in enumerate(KEYS,2)})
            return obj
        report['total']=record(11)
        for row in range(12,sheet.max_row+1):
            name=sheet.cell(row,1).value
            if not name or str(name).startswith('Disclaimer:'):continue
            level=sheet.row_dimensions[row].outlineLevel
            if level==0:
                sector=record(row);sector['groups']=[];report['sectors'].append(sector);group=None
            elif level==2 and sector is not None and sheet.cell(row,11).value is not None:
                group=record(row);group['stocks']=[];sector['groups'].append(group)
            elif sector is not None:
                if group is None:
                    group={'name':'Not Classified','stocks':[],'no_summary':True};sector['groups'].append(group)
                group['stocks'].append(record(row))
        report['row_count']=sum(len(g['stocks']) for s in report['sectors'] for g in s['groups'])
        reports.append(report)
    workbook.close()
    if not reports:raise ValueError('PORT要因分析レポートの見出しが見つかりません。')
    return reports

