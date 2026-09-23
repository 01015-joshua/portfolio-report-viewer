"""Read dated PORT returns/exposures, preserving hierarchy and missing observations."""
from pathlib import Path
from datetime import datetime
import math
import openpyxl

def iso(s):return datetime.strptime(s,'%m/%d/%Y').date().isoformat()
def finite(x):return isinstance(x,(int,float)) and math.isfinite(x)
def extract_analytics(base, start_date, portfolio, benchmark):
    wv=openpyxl.load_workbook(base/'Risk Analysis JIKA.xlsx',data_only=True)
    jpy_path=base/'Risk Analysis JPY.xlsx'
    wr=openpyxl.load_workbook(jpy_path,data_only=True)
    wu=openpyxl.load_workbook(base/'Risk Analysis USD.xlsx',data_only=True)
    v,r,usheet=wv.active,wr.active,wu.active
    if not any(c.value=='Base Currency : USD' for c in usheet[4]):raise ValueError('USDファイルの表示通貨はUSDにしてください。')
    if not all(any(c.value=='Base Currency : JPY' for c in s[4]) for s in (v,r)):raise ValueError('JIKA・JPYファイルの表示通貨はJPYにしてください。')
    def mapping(s):
        result={};universe=None;sector=None
        for n in range(9,s.max_row+1):
            name=s.cell(n,1).value
            if name in (portfolio,benchmark):universe=name;sector=None;key=(universe,'total','')
            elif not name or str(name).startswith('Disclaimer'):continue
            elif s.row_dimensions[n].outlineLevel==1:sector=name;key=(universe,'sector',name)
            elif s.row_dimensions[n].outlineLevel==2:key=(universe,'group',sector+' / '+name)
            else:continue
            if key in result:raise ValueError('分類の階層キーが重複しています。')
            result[key]=n
        return result
    vm,rm,um=mapping(v),mapping(r),mapping(usheet)
    for m in (vm,rm,um):
        if not all((u,'total','') in m for u in (portfolio,benchmark)):raise ValueError('日次ファイルのポートフォリオまたはベンチマークが要因分析レポートと一致していません。')
    if not any(k[1]=='group' for k in rm):raise ValueError('日次レポートに業種グループ（Industry Group）のデータがありません。')
    uc={usheet.cell(8,c).value:c for c in range(2,usheet.max_column+1)}
    if len(uc)!=usheet.max_column-1:raise ValueError('USDの収益区間が重複しています。')
    vc={iso(v.cell(7,c).value):c for c in range(2,v.max_column+1) if v.cell(8,c).value=='EXPOSURE_JIKA_LT'}
    categories={level:sorted({k[2] for k in rm if k[1]==level}) for level in ['sector','group']}
    # Unclassified has no industry-group children: retain it as a separate leaf.
    categories['group'].append('Not Classified')
    def rowkey(u,l,n):return (u,'sector','Not Classified') if n=='Not Classified' else (u,l,n)
    def value(s,m,k,c):
        x=s.cell(m[k],c).value if k in m else None
        return float(x) if finite(x) else None
    daily=[]
    for c in range(2,r.max_column+1):
        start,end=map(iso,r.cell(8,c).value.split('-'))
        if start<start_date:continue
        if start not in vc:raise ValueError('期初の時価総額がありません：'+start)
        if end not in vc:raise ValueError('期末の時価総額がありません：'+end)
        if start>=end:raise ValueError('日次収益の開始日・終了日が不正です。')
        for u in (portfolio,benchmark):
            den=value(v,vm,(u,'total',''),vc[start])
            if den is None or den<=0:raise ValueError('ポートフォリオまたはベンチマークの期初時価総額が不正です：'+start)
        totals={u:value(r,rm,(u,'total',''),c) for u in [portfolio,benchmark]}
        if any(x is None or x<=-100 for x in totals.values()):raise ValueError('合計収益率が不正です：'+end)
        interval=r.cell(8,c).value
        if interval not in uc:raise ValueError('USDの収益区間がありません：'+interval)
        usd_col=uc[interval]
        usd_p=value(usheet,um,(portfolio,'total',''),usd_col)
        usd_b=value(usheet,um,(benchmark,'total',''),usd_col)
        if usd_p is None or usd_b is None or usd_p<=-100 or usd_b<=-100:raise ValueError('USDの合計収益率が不正です：'+interval)
        fx=(1+totals[portfolio]/100)/(1+usd_p/100)
        day={'start':start,'end':end,'p':totals[portfolio]/100,
             'b':(1+usd_b/100)*fx-1,'b_raw':totals[benchmark]/100,'fx_factor':fx,'b_usd':usd_b/100,'p_usd':usd_p/100}
        for level,names in categories.items():
            arr=[]
            for name in names:
                out={'name':name}
                for u,suffix in [(portfolio,'p'),(benchmark,'b')]:
                    k=rowkey(u,level,name);den=value(v,vm,(u,'total',''),vc[start]);mv=value(v,vm,k,vc[start]);ret=value(r,rm,k,c)
                    # Structurally absent benchmark bucket has zero weight, no invented return.
                    out['w'+suffix]=mv/den if mv is not None and den else (0 if k not in vm else None)
                    out['r'+suffix]=ret/100 if ret is not None else None
                    local_usd=value(usheet,um,k,usd_col)
                    out['r'+suffix+'_usd']=local_usd/100 if local_usd is not None else None
                    if suffix=='b':
                        usd_ret=value(usheet,um,k,usd_col)
                        if k in rm and usd_ret is None and ret is not None:raise ValueError('USDベンチマークの分類別収益がありません：'+name+' '+interval)
                        out['rb_raw']=out['rb']
                        out['rb']=(1+usd_ret/100)*fx-1 if usd_ret is not None else None
                arr.append(out)
            day[level]=arr
        day['flow_checks']=[]
        for key,row in vm.items():
            if key[0]!=portfolio or key[1]!='sector':continue
            begin=value(v,vm,key,vc[start]);close=value(v,vm,key,vc[end]);ret=value(r,rm,key,c)
            den=value(v,vm,(portfolio,'total',''),vc[start])
            if begin is not None and close is not None and ret is not None and den:
                day['flow_checks'].append({'sector':key[2],'ratio':(close-begin*(1+ret/100))/den})
        daily.append(day)
    if not daily:raise ValueError('対象年度に日次収益データがありません。')
    if daily[0]['start']!=start_date:raise ValueError('年度の開始日時点のデータがありません。')
    if not all(a['end']==b['start'] for a,b in zip(daily,daily[1:])):raise ValueError('日次収益の区間が連続していないか、日付順に並んでいません。')
    wv.close();wr.close();wu.close()
    return {'daily':daily,'start':daily[0]['start'],'end':daily[-1]['end'],'sources':['Risk Analysis JIKA.xlsx',jpy_path.name,'Risk Analysis USD.xlsx'],'benchmark_method':'USD benchmark returns converted with daily FX factor inferred from portfolio JPY/USD returns'}
